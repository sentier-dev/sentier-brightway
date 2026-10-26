"""BAFU-2026 flows as they land in Brightway, as flat tables for method developers.

Written for the FOODTURE / ADEME request (2026-10-09, "Liste Flux BAFU-Brightway"): one
xlsx with a README, one row per BAFU elementary flow used by the inventory with the EF 3.1
flow it is relinked to (bridge package, factor coverage, Brightway database and code), the
distinct EF 3.1 flows reached, and the process list. Everything is read from a ``data_root``
shaped like the ``~/dds`` checkout through the package's own readers; ``coverage.json`` in
sentier-mappings is the authoritative per-flow sidecar (status, package, tier, reasons) and
the bridge packages supply the EF target. The two are cross-checked and any disagreement
raises ``ValueError`` rather than producing a silently wrong table.
"""

from __future__ import annotations

import json
import re
import subprocess
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Mapping

import pandas as pd

from ._frames import data_root_error
from ._version import __version__
from .bridge import BridgeEntry, load_bridge
from .constants import (
    BIOSPHERE_DB,
    BRIDGE_FOLDER,
    DEFAULT_SOURCE,
    REPO_INVENTORY,
    REPO_MAPPINGS,
    REPO_METHODS,
    REPO_VOCAB,
    citation,
    inventory_db,
    residual_db,
)
from .flows import load_ef_flows, load_source_flows
from .inventory import Inventory, load_inventory
from .methods import MethodSpec, load_methods

COVERAGE_FILE = "coverage.json"
STATUSES = frozenset({"mapped", "unmapped"})
_PACKAGE_PREFIX = re.compile(r"^biosphere-\d+-")

COVERAGE_COLUMNS = (
    "bafu_code",
    "bafu_name",
    "bafu_unit",
    "bafu_compartment",
    "bafu_sub_compartment",
    "status",
    "package",
    "tier",
    "placement",
    "reason",
    "detail",
    "caveats",
    "cas_conflict",
)


def _cell(value: object) -> str:
    """One coverage field as a spreadsheet cell: lists joined, dicts as JSON, None empty."""
    if value is None:
        return ""
    if isinstance(value, list):
        return "; ".join(str(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _coverage_path(data_root: Path) -> Path:
    return Path(data_root) / REPO_MAPPINGS / "data" / BRIDGE_FOLDER / COVERAGE_FILE


def load_coverage(data_root: Path) -> pd.DataFrame:
    """``coverage.json`` flattened to one row per BAFU flow, columns ``COVERAGE_COLUMNS``."""
    path = _coverage_path(data_root)
    if not path.is_file():
        raise data_root_error(COVERAGE_FILE, path)
    rows = json.loads(path.read_text()).get("coverage")
    if not isinstance(rows, list):
        raise ValueError(f"{path}: expected a 'coverage' list")
    out = []
    for i, row in enumerate(rows):
        source = row.get("source") or {}
        code = source.get("code")
        if not code:
            raise ValueError(f"{path} row {i}: missing source.code")
        status = row.get("status")
        if status not in STATUSES:
            raise ValueError(f"{path} row {i} ({code}): unknown status {status!r}")
        context = list(source.get("context") or [])
        out.append(
            {
                "bafu_code": code,
                "bafu_name": _cell(source.get("name")),
                "bafu_unit": _cell(source.get("unit")),
                "bafu_compartment": context[0] if context else "",
                "bafu_sub_compartment": context[1] if len(context) > 1 else "",
                "status": status,
                "package": _PACKAGE_PREFIX.sub("", _cell(row.get("package"))),
                "tier": _cell(row.get("tier")),
                "placement": _cell(row.get("placement")),
                "reason": _cell(row.get("reason")),
                "detail": _cell(row.get("detail")),
                "caveats": _cell(row.get("caveats")),
                "cas_conflict": _cell(row.get("cas_conflict")),
            }
        )
    return pd.DataFrame(out, columns=list(COVERAGE_COLUMNS))


# ------------------------------------------------------------------ EF side lookups


def ef_flow_table(ef_flows: pd.DataFrame, methods: tuple[MethodSpec, ...]) -> pd.DataFrame:
    """Every EF 3.1 flow the Brightway biosphere can hold: columns ``code, name, context, cas``.

    Same precedence as ``build.build_biosphere``: the vocab names a flow first, a method's
    CF table names flows the vocab lacks, and the CF context (three levels) beats the vocab
    compartment because it is more specific."""
    names: dict[str, str] = dict(zip(ef_flows["code"], ef_flows["name"]))
    cas: dict[str, str] = {c: (v or "") for c, v in zip(ef_flows["code"], ef_flows["cas_number"])}
    context: dict[str, str] = {
        c: " / ".join(cats) for c, cats in zip(ef_flows["code"], ef_flows["categories"])
    }
    cf_context: dict[str, str] = {}
    for spec in methods:
        for code, name in spec.flow_name.items():
            names.setdefault(code, name)
        for code, ctx in spec.flow_context.items():
            if ctx:
                cf_context.setdefault(code, " / ".join(ctx))
    context.update(cf_context)
    codes = sorted(names)
    return pd.DataFrame(
        {
            "code": codes,
            "name": [names[c] for c in codes],
            "context": [context.get(c, "") for c in codes],
            "cas": [cas.get(c, "") for c in codes],
        }
    )


def methods_per_flow(methods: tuple[MethodSpec, ...]) -> dict[str, tuple[str, ...]]:
    """``{ef_flow_code: (method_id, ...)}`` over the global characterization factors."""
    out: dict[str, list[str]] = {}
    for spec in methods:
        for code, _factor in spec.cfs:
            ids = out.setdefault(code, [])
            if spec.method_id not in ids:
                ids.append(spec.method_id)
    return {code: tuple(ids) for code, ids in out.items()}


# ------------------------------------------------------------------ bafu_to_ef sheet

BAFU_TO_EF_COLUMNS = (
    "bafu_name",
    "bafu_code",
    "bafu_unit",
    "bafu_compartment",
    "bafu_sub_compartment",
    "bafu_cas",
    "exchange_count",
    "status",
    "package",
    "tier",
    "placement",
    "ef_uuid",
    "ef_name",
    "ef_context",
    "ef_unit",
    "ef_cas",
    "conversion_factor",
    "ef_methods_with_factor",
    "brightway_database",
    "brightway_code",
    "reason",
    "detail",
    "caveats",
    "cas_conflict",
)


def _first(codes) -> str:
    return ", ".join(sorted(codes)[:5])


def _check_coverage_against_inputs(
    coverage: pd.DataFrame,
    bridge: Mapping[str, BridgeEntry],
    ef_table: pd.DataFrame,
    exchanges: pd.DataFrame,
) -> None:
    """Every disagreement between the sidecar, the bridge and the inventory is an error."""
    dup = coverage["bafu_code"][coverage["bafu_code"].duplicated()]
    if not dup.empty:
        raise ValueError(f"coverage.json: duplicate bafu_code: {_first(dup)}")
    used = set(exchanges.loc[exchanges["flow_type"] == "biosphere", "flow"])
    covered = set(coverage["bafu_code"])
    if used != covered:
        raise ValueError(
            "coverage.json rows differ from the flows the inventory uses: "
            f"used but not covered: {_first(used - covered) or '-'}; "
            f"covered but not used: {_first(covered - used) or '-'}"
        )
    mapped = set(coverage.loc[coverage["status"] == "mapped", "bafu_code"])
    unmapped = covered - mapped
    missing = mapped - set(bridge)
    if missing:
        raise ValueError(f"coverage says mapped but no bridge entry: {_first(missing)}")
    stray = unmapped & set(bridge)
    if stray:
        raise ValueError(f"coverage says unmapped but the bridge maps it: {_first(stray)}")
    known = set(ef_table["code"])
    unknown = {bridge[c].target_code for c in mapped} - known
    if unknown:
        raise ValueError(f"bridge targets unknown to vocab and EF methods: {_first(unknown)}")


def build_bafu_to_ef(
    coverage: pd.DataFrame,
    bridge: Mapping[str, BridgeEntry],
    ef_table: pd.DataFrame,
    bafu_flows: pd.DataFrame,
    exchanges: pd.DataFrame,
    methods: tuple[MethodSpec, ...],
    source: str = DEFAULT_SOURCE,
) -> pd.DataFrame:
    """One row per BAFU flow the inventory uses, with the EF 3.1 flow it lands on."""
    _check_coverage_against_inputs(coverage, bridge, ef_table, exchanges)
    bio = exchanges[exchanges["flow_type"] == "biosphere"]
    counts = bio["flow"].value_counts()
    bafu_cas = {c: (v or "") for c, v in zip(bafu_flows["code"], bafu_flows["cas_number"])}
    ef = ef_table.set_index("code")
    per_flow = methods_per_flow(methods)
    rows = []
    for rec in coverage.to_dict("records"):
        code = rec["bafu_code"]
        row = dict(rec)
        row["bafu_cas"] = bafu_cas.get(code, "")
        row["exchange_count"] = int(counts.get(code, 0))
        if rec["status"] == "mapped":
            entry = bridge[code]
            target = ef.loc[entry.target_code]
            row.update(
                ef_uuid=entry.target_code,
                ef_name=target["name"],
                ef_context=target["context"],
                ef_unit=entry.target_unit,
                ef_cas=target["cas"],
                conversion_factor=entry.conversion_factor,
                ef_methods_with_factor=len(per_flow.get(entry.target_code, ())),
                brightway_database=BIOSPHERE_DB,
                brightway_code=entry.target_code,
            )
        else:
            row.update(
                ef_uuid="",
                ef_name="",
                ef_context="",
                ef_unit="",
                ef_cas="",
                conversion_factor=None,
                ef_methods_with_factor=0,
                brightway_database=residual_db(source),
                brightway_code=code,
            )
        rows.append(row)
    df = pd.DataFrame(rows, columns=list(BAFU_TO_EF_COLUMNS))
    return df.sort_values(
        ["bafu_compartment", "bafu_sub_compartment", "bafu_name"], kind="stable"
    ).reset_index(drop=True)


# ------------------------------------------------------------------ ef_flows, processes

EF_FLOWS_COLUMNS = (
    "ef_uuid",
    "ef_name",
    "ef_context",
    "ef_unit",
    "ef_cas",
    "bafu_source_count",
    "exchange_count",
    "ef_methods_with_factor",
    "ef_methods",
)

PROCESS_COLUMNS = (
    "name",
    "reference_product",
    "unit",
    "location",
    "sector",
    "process_type",
    "code",
    "brightway_database",
)


def build_ef_flows(
    bafu_to_ef: pd.DataFrame,
    bridge: Mapping[str, BridgeEntry],
    methods: tuple[MethodSpec, ...],
) -> pd.DataFrame:
    """The distinct EF 3.1 flows the mapped BAFU flows land on."""
    units: dict[str, str] = {}
    for entry in bridge.values():  # package order, first wins: same rule as build()
        units.setdefault(entry.target_code, entry.target_unit)
    per_flow = methods_per_flow(methods)
    mapped = bafu_to_ef[bafu_to_ef["status"] == "mapped"]
    grouped = mapped.groupby("ef_uuid", sort=False).agg(
        ef_name=("ef_name", "first"),
        ef_context=("ef_context", "first"),
        ef_cas=("ef_cas", "first"),
        bafu_source_count=("bafu_code", "size"),
        exchange_count=("exchange_count", "sum"),
    )
    grouped = grouped.reset_index()
    grouped["ef_unit"] = grouped["ef_uuid"].map(units)
    grouped["ef_methods"] = grouped["ef_uuid"].map(lambda c: "; ".join(per_flow.get(c, ())))
    grouped["ef_methods_with_factor"] = grouped["ef_uuid"].map(lambda c: len(per_flow.get(c, ())))
    df = grouped[list(EF_FLOWS_COLUMNS)]
    return df.sort_values(["ef_context", "ef_name"], kind="stable").reset_index(drop=True)


def build_processes(inventory: Inventory) -> pd.DataFrame:
    """The technosphere nomenclature: one row per process, as installed in Brightway."""
    p = inventory.processes
    dup = p["process_id"][p["process_id"].duplicated()]
    if not dup.empty:
        raise ValueError(f"processes: duplicate process_id: {_first(dup)}")
    df = pd.DataFrame(
        {
            "name": p["name"],
            "reference_product": p["reference_product"],
            "unit": p["reference_unit"],
            "location": p["location"],
            "sector": p["sector"] if "sector" in p.columns else "",
            "process_type": p["process_type"],
            "code": p["process_id"],
            "brightway_database": inventory_db(inventory.source),
        },
        columns=list(PROCESS_COLUMNS),
    )
    return df.sort_values(["sector", "name", "location"], kind="stable").reset_index(drop=True)


# ------------------------------------------------------------------ provenance

REPOS = (REPO_INVENTORY, REPO_VOCAB, REPO_METHODS, REPO_MAPPINGS)


@dataclass(frozen=True)
class RepoPin:
    repo: str
    sha: str  # full git SHA of the checkout read, or "unknown"
    dirty: bool  # working tree had uncommitted or untracked changes

    @property
    def label(self) -> str:
        return f"{self.sha}+dirty" if self.dirty else self.sha


def _git_output(repo_dir: Path, *args: str) -> str | None:
    try:
        done = subprocess.run(
            ["git", "-C", str(repo_dir), *args], capture_output=True, text=True, check=False
        )
    except OSError:
        return None
    return done.stdout.strip() if done.returncode == 0 else None


def repo_pins(data_root: Path) -> tuple[RepoPin, ...]:
    """The git state of each data-repo checkout under ``data_root``; never raises."""
    pins = []
    for repo in REPOS:
        repo_dir = Path(data_root) / repo
        sha = _git_output(repo_dir, "rev-parse", "HEAD") if (repo_dir / ".git").exists() else None
        status = _git_output(repo_dir, "status", "--porcelain") if sha else None
        pins.append(RepoPin(repo=repo, sha=sha or "unknown", dirty=bool(status)))
    return tuple(pins)


# ------------------------------------------------------------------ README and outputs

SHEETS = ("README", "bafu_to_ef", "ef_flows", "processes")
MAX_COLUMN_WIDTH = 60
EF_REFERENCE = (
    "EF 3.1 elementary flows and characterization factors: European Commission, Joint "
    "Research Centre, Environmental Footprint reference package 3.1 (as delivered by "
    "sentier-methods and sentier-vocab)."
)

COLUMN_NOTES = {
    "bafu_to_ef": (
        "One row per BAFU-2026 elementary flow used by at least one process. bafu_* columns "
        "are the BAFU flow (code = BAFU-2026 v1 flow id, context as BAFU spells it). "
        "exchange_count = biosphere exchange rows in the inventory using this flow. status = "
        "mapped (relinked to an EF 3.1 flow) or unmapped (kept under its BAFU name). package = "
        "which bridge package of sentier-mappings placed it: curated, inferred, matched, or "
        "nomenclature (valid EF flow key, but no factor in any EF 3.1 method). ef_* columns are "
        "the EF 3.1 target (ef_uuid = EF flow UUID, the Brightway code); conversion_factor "
        "multiplies the BAFU amount. ef_methods_with_factor = number of the 25 EF 3.1 methods "
        "with a global factor for the target (0 for nomenclature and unmapped rows). "
        "brightway_database / brightway_code = where the flow lives after import. reason and "
        "detail explain unmapped rows; caveats and cas_conflict carry the curators' notes."
    ),
    "ef_flows": (
        "The distinct EF 3.1 flows the mapped BAFU flows land on: the biosphere nomenclature "
        "a BAFU-based inventory actually carries in Brightway. bafu_source_count = BAFU flows "
        "relinked onto this EF flow; exchange_count = their summed exchange rows; ef_methods = "
        "EF 3.1 methods with a global factor for it."
    ),
    "processes": (
        "The technosphere nomenclature: one row per BAFU-2026 process as installed in the "
        "Brightway database named in brightway_database; code is the process id."
    ),
}


@dataclass(frozen=True)
class ExportResult:
    xlsx: Path
    counts: dict[str, int]
    headline: dict[str, int]


def headline_numbers(
    bafu_to_ef: pd.DataFrame, ef_flows: pd.DataFrame, processes: pd.DataFrame
) -> dict[str, int]:
    mapped = bafu_to_ef["status"] == "mapped"
    return {
        "flows_used": int(len(bafu_to_ef)),
        "flows_mapped": int(mapped.sum()),
        "flows_unmapped": int((~mapped).sum()),
        "flows_nomenclature": int((bafu_to_ef["package"] == "nomenclature").sum()),
        "ef_flows_reached": int(len(ef_flows)),
        "processes": int(len(processes)),
    }


def readme_rows(
    pins: tuple[RepoPin, ...],
    counts: Mapping[str, int],
    headline: Mapping[str, int],
    today: date,
    source: str = DEFAULT_SOURCE,
) -> list[tuple[str, str]]:
    """The README sheet as (key, value) rows; a blank key starts a free-text paragraph."""
    rows: list[tuple[str, str]] = [
        ("title", f"{source} elementary flows and processes as implemented in Brightway"),
        ("generated", today.isoformat()),
        ("generator", f"sentier-brightway {__version__}, scripts/export_flow_lists.py"),
        ("inventory folder", f"{REPO_INVENTORY}/data/{BRIDGE_FOLDER.split('__')[0]}"),
        ("bridge folder", f"{REPO_MAPPINGS}/data/{BRIDGE_FOLDER}"),
    ]
    rows += [(f"git {p.repo}", p.label) for p in pins]
    rows += [(f"rows {name}", str(n)) for name, n in counts.items()]
    rows += [(key, str(n)) for key, n in headline.items()]
    rows += [("citation", citation(source)), ("ef reference", EF_REFERENCE)]
    rows.append(
        (
            "",
            "How to read this: in our Brightway import the BAFU elementary flows are not kept "
            "under their own nomenclature. Each mapped flow is relinked to an EF 3.1 flow in "
            f"the database '{BIOSPHERE_DB}', keyed by the EF flow UUID, so a method aligned "
            "with EF 3.1 flow keys characterises these inventories directly. The unmapped "
            f"flows stay under their BAFU name and code in the database '{residual_db(source)}' "
            "and receive no factor. Rows placed by the 'nomenclature' package point at a valid "
            "EF 3.1 flow that no EF 3.1 method characterises; a new method may add factors "
            "for them.",
        )
    )
    for name in SHEETS[1:]:
        rows.append(("", f"Sheet {name}: {COLUMN_NOTES[name]}"))
    return rows


def readme_markdown(rows: list[tuple[str, str]]) -> str:
    lines = ["# BAFU-2026 flow lists as implemented in Brightway", ""]
    for key, value in rows:
        lines.append(f"- {key}: {value}" if key else f"\n{value}\n")
    return "\n".join(lines).rstrip() + "\n"


def _set_widths(ws) -> None:
    for column in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)):
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(MAX_COLUMN_WIDTH, width + 2)


def _verify_workbook(xlsx: Path, frames: Mapping[str, pd.DataFrame]) -> None:
    import openpyxl  # a dependency already; imported here so the module stays light

    wb = openpyxl.load_workbook(xlsx, read_only=True)
    try:
        if wb.sheetnames != list(SHEETS):
            raise ValueError(f"{xlsx}: sheets {wb.sheetnames} != {list(SHEETS)}")
        for name, df in frames.items():
            rows = wb[name].max_row - 1
            if rows != len(df):
                raise ValueError(f"{xlsx}: sheet {name} has {rows} rows, expected {len(df)}")
    finally:
        wb.close()


def write_outputs(
    frames: Mapping[str, pd.DataFrame],
    readme: list[tuple[str, str]],
    out_dir: Path,
    stem: str,
    overwrite: bool = False,
) -> Path:
    """CSVs under ``out_dir/csv``, ``README.md``, then the xlsx; returns the xlsx path."""
    out_dir = Path(out_dir)
    xlsx = out_dir / f"{stem}.xlsx"
    if xlsx.exists() and not overwrite:
        raise FileExistsError(f"{xlsx} exists; pass overwrite=True (--overwrite) to replace it")
    (out_dir / "csv").mkdir(parents=True, exist_ok=True)
    for name, df in frames.items():
        df.to_csv(out_dir / "csv" / f"{name}.csv", index=False, encoding="utf-8")
    (out_dir / "README.md").write_text(readme_markdown(readme), encoding="utf-8")
    readme_df = pd.DataFrame(readme, columns=["key", "value"])
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        readme_df.to_excel(writer, sheet_name="README", index=False)
        for name, df in frames.items():
            df.to_excel(writer, sheet_name=name, index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            if ws.title != "README":
                ws.auto_filter.ref = ws.dimensions
            _set_widths(ws)
    _verify_workbook(xlsx, frames)
    return xlsx


def export(
    data_root: Path,
    out_dir: Path,
    today: date | None = None,
    overwrite: bool = False,
    source: str = DEFAULT_SOURCE,
) -> ExportResult:
    """Read the four data repos under ``data_root`` and write the workbook into ``out_dir``."""
    today = today or date.today()
    inventory = load_inventory(data_root, source=source)
    methods = load_methods(data_root)
    bridge = load_bridge(data_root)
    ef_table = ef_flow_table(load_ef_flows(data_root), methods)
    bafu_to_ef = build_bafu_to_ef(
        coverage=load_coverage(data_root),
        bridge=bridge,
        ef_table=ef_table,
        bafu_flows=load_source_flows(data_root, source),
        exchanges=inventory.exchanges,
        methods=methods,
        source=inventory.source,
    )
    ef_flows = build_ef_flows(bafu_to_ef, bridge, methods)
    processes = build_processes(inventory)
    frames = {"bafu_to_ef": bafu_to_ef, "ef_flows": ef_flows, "processes": processes}
    counts = {name: int(len(df)) for name, df in frames.items()}
    headline = headline_numbers(bafu_to_ef, ef_flows, processes)
    readme = readme_rows(repo_pins(data_root), counts, headline, today, inventory.source)
    stem = f"{BRIDGE_FOLDER}_flow-lists_{today.isoformat()}"
    xlsx = write_outputs(frames, readme, out_dir, stem, overwrite=overwrite)
    return ExportResult(xlsx=xlsx, counts=counts, headline=headline)
