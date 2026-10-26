"""Provenance, README text and file writing for the flow-lists workbook.

Split out of ``flow_lists`` so that module stays focused on building the tables; this one
knows the sheet order, the README wording, and how the xlsx and its side artifacts are
written and verified.
"""

from __future__ import annotations

import subprocess
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Mapping

import pandas as pd

from ._version import __version__
from .constants import (
    BIOSPHERE_DB,
    BRIDGE_FOLDER,
    DEFAULT_SOURCE,
    REPO_INVENTORY,
    REPO_MAPPINGS,
    REPO_METHODS,
    REPO_VOCAB,
    citation,
    residual_db,
)

# ------------------------------------------------------------------ provenance

REPOS = (REPO_INVENTORY, REPO_VOCAB, REPO_METHODS, REPO_MAPPINGS)


@dataclass(frozen=True)
class RepoPin:
    repo: str
    sha: str  # full git SHA of the checkout read, or "unknown"
    dirty: bool  # the data/ folder had uncommitted or untracked changes

    @property
    def label(self) -> str:
        return f"{self.sha}+dirty" if self.dirty else self.sha


def _git_output(repo_dir: Path, *args: str) -> str | None:
    try:
        done = subprocess.run(
            ["git", "-C", str(repo_dir), *args], capture_output=True, text=True, check=False
        )
    except OSError:
        return None
    return done.stdout.strip() if done.returncode == 0 else None


def repo_pins(data_root: Path) -> tuple[RepoPin, ...]:
    """The git state of each data-repo checkout under ``data_root``; never raises.

    ``dirty`` looks only at the ``data/`` folder, the part of the repo that is read; an
    edited ``.gitignore`` or README elsewhere does not taint the pin."""
    pins = []
    for repo in REPOS:
        repo_dir = Path(data_root) / repo
        sha = _git_output(repo_dir, "rev-parse", "HEAD") if (repo_dir / ".git").exists() else None
        status = _git_output(repo_dir, "status", "--porcelain", "--", "data") if sha else None
        pins.append(RepoPin(repo=repo, sha=sha or "unknown", dirty=bool(status)))
    return tuple(pins)


# ------------------------------------------------------------------ README and outputs

SHEETS = ("README", "bafu_to_ef", "ef_flows", "processes")
MAX_COLUMN_WIDTH = 60
EF_REFERENCE = (
    "EF 3.1 elementary flows and characterization factors: European Commission, Joint "
    "Research Centre, Environmental Footprint reference package 3.1 (as delivered by "
    "sentier-methods and sentier-vocab)."
)

COLUMN_NOTES = {
    "bafu_to_ef": (
        "One row per BAFU-2026 elementary flow used by at least one process. bafu_* columns "
        "are the BAFU flow (code = BAFU-2026 v1 flow id, context as BAFU spells it). "
        "exchange_count = biosphere exchange rows in the inventory using this flow. status = "
        "mapped (relinked to an EF 3.1 flow) or unmapped (kept under its BAFU name). package = "
        "which bridge package of sentier-mappings placed it: curated, inferred, matched, or "
        "nomenclature (valid EF flow key, but no factor in any EF 3.1 method). ef_* columns are "
        "the EF 3.1 target (ef_uuid = EF flow UUID, the Brightway code); conversion_factor "
        "multiplies the BAFU amount. ef_methods_with_factor = number of the 25 EF 3.1 methods "
        "with a global factor for the target (0 for nomenclature and unmapped rows). "
        "brightway_database / brightway_code = where the flow lives after import. reason and "
        "detail explain unmapped rows; caveats and cas_conflict carry the curators' notes."
    ),
    "ef_flows": (
        "The distinct EF 3.1 flows the mapped BAFU flows land on: the biosphere nomenclature "
        "a BAFU-based inventory actually carries in Brightway. bafu_source_count = BAFU flows "
        "relinked onto this EF flow; exchange_count = their summed exchange rows; ef_methods = "
        "EF 3.1 methods with a global factor for it."
    ),
    "processes": (
        "The technosphere nomenclature: one row per BAFU-2026 process as installed in the "
        "Brightway database named in brightway_database; code is the process id."
    ),
}


@dataclass(frozen=True)
class ExportResult:
    xlsx: Path
    counts: dict[str, int]
    headline: dict[str, int]


def headline_numbers(
    bafu_to_ef: pd.DataFrame, ef_flows: pd.DataFrame, processes: pd.DataFrame
) -> dict[str, int]:
    mapped = bafu_to_ef["status"] == "mapped"
    return {
        "flows_used": int(len(bafu_to_ef)),
        "flows_mapped": int(mapped.sum()),
        "flows_unmapped": int((~mapped).sum()),
        "flows_nomenclature": int((bafu_to_ef["package"] == "nomenclature").sum()),
        "ef_flows_reached": int(len(ef_flows)),
        "processes": int(len(processes)),
    }


def readme_rows(
    pins: tuple[RepoPin, ...],
    counts: Mapping[str, int],
    headline: Mapping[str, int],
    today: date,
    source: str = DEFAULT_SOURCE,
) -> list[tuple[str, str]]:
    """The README sheet as (key, value) rows; a blank key starts a free-text paragraph."""
    rows: list[tuple[str, str]] = [
        ("title", f"{source} elementary flows and processes as implemented in Brightway"),
        ("generated", today.isoformat()),
        ("generator", f"sentier-brightway {__version__}, scripts/export_flow_lists.py"),
        ("inventory folder", f"{REPO_INVENTORY}/data/{BRIDGE_FOLDER.split('__')[0]}"),
        ("bridge folder", f"{REPO_MAPPINGS}/data/{BRIDGE_FOLDER}"),
    ]
    rows += [(f"git {p.repo}", p.label) for p in pins]
    rows += [(f"rows {name}", str(n)) for name, n in counts.items()]
    rows += [(key, str(n)) for key, n in headline.items()]
    rows += [("citation", citation(source)), ("ef reference", EF_REFERENCE)]
    rows.append(
        (
            "",
            "How to read this: in our Brightway import the BAFU elementary flows are not kept "
            "under their own nomenclature. Each mapped flow is relinked to an EF 3.1 flow in "
            f"the database '{BIOSPHERE_DB}', keyed by the EF flow UUID, so a method aligned "
            "with EF 3.1 flow keys characterises these inventories directly. The unmapped "
            f"flows stay under their BAFU name and code in the database '{residual_db(source)}' "
            "and receive no factor. Rows placed by the 'nomenclature' package point at a valid "
            "EF 3.1 flow that no EF 3.1 method characterises; a new method may add factors "
            "for them.",
        )
    )
    for name in SHEETS[1:]:
        rows.append(("", f"Sheet {name}: {COLUMN_NOTES[name]}"))
    return rows


def readme_markdown(rows: list[tuple[str, str]]) -> str:
    lines = ["# BAFU-2026 flow lists as implemented in Brightway", ""]
    for key, value in rows:
        lines.append(f"- {key}: {value}" if key else f"\n{value}\n")
    return "\n".join(lines).rstrip() + "\n"


def _set_widths(ws) -> None:
    for column in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)):
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(MAX_COLUMN_WIDTH, width + 2)


def _verify_workbook(xlsx: Path, frames: Mapping[str, pd.DataFrame]) -> None:
    import openpyxl  # a dependency already; imported here so the module stays light

    wb = openpyxl.load_workbook(xlsx, read_only=True)
    try:
        if wb.sheetnames != list(SHEETS):
            raise ValueError(f"{xlsx}: sheets {wb.sheetnames} != {list(SHEETS)}")
        for name, df in frames.items():
            rows = wb[name].max_row - 1
            if rows != len(df):
                raise ValueError(f"{xlsx}: sheet {name} has {rows} rows, expected {len(df)}")
    finally:
        wb.close()


def write_outputs(
    frames: Mapping[str, pd.DataFrame],
    readme: list[tuple[str, str]],
    out_dir: Path,
    stem: str,
    overwrite: bool = False,
) -> Path:
    """CSVs under ``out_dir/csv``, ``README.md``, then the xlsx; returns the xlsx path."""
    out_dir = Path(out_dir)
    xlsx = out_dir / f"{stem}.xlsx"
    if xlsx.exists() and not overwrite:
        raise FileExistsError(f"{xlsx} exists; pass overwrite=True (--overwrite) to replace it")
    (out_dir / "csv").mkdir(parents=True, exist_ok=True)
    for name, df in frames.items():
        df.to_csv(out_dir / "csv" / f"{name}.csv", index=False, encoding="utf-8")
    (out_dir / "README.md").write_text(readme_markdown(readme), encoding="utf-8")
    readme_df = pd.DataFrame(readme, columns=["key", "value"])
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        readme_df.to_excel(writer, sheet_name="README", index=False)
        for name, df in frames.items():
            df.to_excel(writer, sheet_name=name, index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            if ws.title != "README":
                ws.auto_filter.ref = ws.dimensions
            _set_widths(ws)
    _verify_workbook(xlsx, frames)
    return xlsx
