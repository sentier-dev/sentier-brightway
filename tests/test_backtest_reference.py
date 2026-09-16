import zipfile

import pytest

from sentier_brightway.backtest.reference import BafuReference, demojibake, split_product
from tests.conftest import REF_P1_CLIMATE_PER_MJ


def test_split_product_on_last_separator():
    assert split_product("Electricity, low voltage, at grid - CH") == (
        "Electricity, low voltage, at grid",
        "CH",
    )
    assert split_product("Cotton fibres, at farm - US") == ("Cotton fibres, at farm", "US")
    assert split_product("Some - thing - RER") == ("Some - thing", "RER")


def test_demojibake_roundtrip():
    assert demojibake("WÃƒÂ¤rme") == "Wärme"
    assert demojibake("plain") == "plain"


def test_reads_headers_units_and_scores(bafu_xlsx):
    ref = BafuReference.from_path(bafu_xlsx)
    assert list(ref.frame.columns[:3]) == ["name", "location", "unit"]
    assert len(ref.frame) == 3
    row = ref.frame.set_index(["name", "location"]).loc[
        ("Electricity, low voltage, at grid", "CH")
    ]
    assert row["unit"] == "MJ"
    assert row["climate"] == pytest.approx(REF_P1_CLIMATE_PER_MJ)
    assert row["acid"] == 0.0
    assert row["water"] != row["water"]  # NaN for blank cells
    assert ref.blank_cells > 0
    assert "Wärme, ab Kessel" in set(ref.frame["name"])


def test_zip_input_is_accepted(bafu_xlsx, tmp_path):
    z = tmp_path / "ref.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(bafu_xlsx, "BAFU-2026 v1_LCIA Results/x.xlsx")
    assert len(BafuReference.from_path(z).frame) == 3


def test_duplicate_product_is_an_error(bafu_xlsx):
    import openpyxl

    wb = openpyxl.load_workbook(bafu_xlsx)
    ws = wb.active
    ws.append([c.value for c in ws[3]])
    wb.save(bafu_xlsx)
    with pytest.raises(ValueError, match="duplicate"):
        BafuReference.from_path(bafu_xlsx)


def test_missing_ef_header_is_an_error(bafu_xlsx):
    import openpyxl

    wb = openpyxl.load_workbook(bafu_xlsx)
    ws = wb.active
    ws.cell(row=2, column=6).value = "Not an EF column"
    wb.save(bafu_xlsx)
    with pytest.raises(ValueError, match="Climate change"):
        BafuReference.from_path(bafu_xlsx)
